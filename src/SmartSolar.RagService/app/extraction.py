"""Text extraction for PDF, DOCX, TXT and Markdown with resource limits.

Scanned PDFs (almost no extractable text) are reported as OCR_REQUIRED instead of being indexed
empty; OCR is intentionally not performed by this service.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional

from .config import DOCX, MARKDOWN, PDF, TEXT

READY = "READY"
OCR_REQUIRED = "OCR_REQUIRED"
UNSUPPORTED_CONTENT = "UNSUPPORTED_CONTENT"


@dataclass
class Section:
    text: str
    page: Optional[int] = None
    heading: Optional[str] = None


@dataclass
class Extraction:
    status: str
    sections: List[Section] = field(default_factory=list)
    page_count: Optional[int] = None
    characters: int = 0
    error: Optional[str] = None


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SPACES = re.compile(r"[ \t ]+")


def normalize(text: str) -> str:
    text = _CONTROL.sub(" ", text.replace("\r\n", "\n").replace("\r", "\n"))
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract(data: bytes, content_type: str, max_pages: int, max_characters: int, max_uncompressed_bytes: int) -> Extraction:
    try:
        if content_type == PDF:
            return _pdf(data, max_pages, max_characters)
        if content_type == DOCX:
            return _docx(data, max_characters, max_uncompressed_bytes)
        if content_type in (TEXT, MARKDOWN):
            return _text(data, content_type == MARKDOWN, max_characters)
    except Exception as exc:  # noqa: BLE001 - malformed input must never crash the worker
        return Extraction(UNSUPPORTED_CONTENT, error=f"unreadable document ({type(exc).__name__})")
    return Extraction(UNSUPPORTED_CONTENT, error="unsupported content type")


def _pdf(data: bytes, max_pages: int, max_characters: int) -> Extraction:
    if not data.startswith(b"%PDF-"):
        return Extraction(UNSUPPORTED_CONTENT, error="not a PDF")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        return Extraction(UNSUPPORTED_CONTENT, error="encrypted PDF")
    pages = len(reader.pages)
    if pages > max_pages:
        return Extraction(UNSUPPORTED_CONTENT, page_count=pages, error=f"more than {max_pages} pages")
    sections: List[Section] = []
    total = 0
    for number, page in enumerate(reader.pages, start=1):
        text = normalize(page.extract_text() or "")
        total += len(text)
        if total > max_characters:
            return Extraction(UNSUPPORTED_CONTENT, page_count=pages, error="document text exceeds the size limit")
        if text:
            sections.append(Section(text, page=number))
    # Fewer than ~50 characters per page on average: a scan or image-only PDF.
    if total < max(200, 50 * pages):
        return Extraction(OCR_REQUIRED, page_count=pages, characters=total, error="little or no extractable text; OCR is required")
    return Extraction(READY, sections, page_count=pages, characters=total)


def _docx(data: bytes, max_characters: int, max_uncompressed_bytes: int) -> Extraction:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > max_uncompressed_bytes:
            return Extraction(UNSUPPORTED_CONTENT, error="archive expands beyond the size limit")
        if "word/document.xml" not in archive.namelist():
            return Extraction(UNSUPPORTED_CONTENT, error="not a Word document")
    import docx

    document = docx.Document(io.BytesIO(data))
    sections: List[Section] = []
    heading: Optional[str] = None
    buffer: List[str] = []
    total = 0

    def flush() -> None:
        if buffer:
            sections.append(Section(normalize("\n".join(buffer)), heading=heading))
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        total += len(text)
        if total > max_characters:
            return Extraction(UNSUPPORTED_CONTENT, error="document text exceeds the size limit")
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading") or style.lower() == "title":
            flush()
            heading = text[:200]
        buffer.append(text)
    flush()
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = normalize("\n".join(r for r in rows if r.strip(" |")))
        total += len(text)
        if total > max_characters:
            return Extraction(UNSUPPORTED_CONTENT, error="document text exceeds the size limit")
        if text:
            sections.append(Section(text, heading=heading))
    if total < 20:
        return Extraction(UNSUPPORTED_CONTENT, characters=total, error="no extractable text")
    return Extraction(READY, [s for s in sections if s.text], characters=total)


def _text(data: bytes, markdown: bool, max_characters: int) -> Extraction:
    try:
        raw = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return Extraction(UNSUPPORTED_CONTENT, error="text is not valid UTF-8")
    if len(raw) > max_characters:
        return Extraction(UNSUPPORTED_CONTENT, error="document text exceeds the size limit")
    text = normalize(raw)
    if len(text) < 20:
        return Extraction(UNSUPPORTED_CONTENT, characters=len(text), error="no extractable text")
    if not markdown:
        return Extraction(READY, [Section(text)], characters=len(text))
    sections: List[Section] = []
    heading: Optional[str] = None
    buffer: List[str] = []
    for line in text.split("\n"):
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if match:
            if buffer:
                sections.append(Section("\n".join(buffer).strip(), heading=heading))
                buffer = []
            heading = match.group(2).strip()[:200]
        buffer.append(line)
    if buffer:
        sections.append(Section("\n".join(buffer).strip(), heading=heading))
    return Extraction(READY, [s for s in sections if s.text], characters=len(text))
